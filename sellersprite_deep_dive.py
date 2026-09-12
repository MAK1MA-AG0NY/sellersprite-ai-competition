# -*- coding: utf-8 -*-
"""
卖家精灵深度开发：历史报告 → 竞品ASIN提取 → 竞品详情批量查询 → Excel导出
=========================================================================
流程（一条龙模式）：
  1) POST /v3/api/ai-analysis/usage-log      分页拉取历史任务列表
  2) GET  /v3/api/ai-analysis/task/{taskId}  拉任务详情，content JSON 里提取竞品 ASIN
  3) 竞品查询 /v3/competitor-lookup：
     - 首次运行自动"探测"页面真实调用的接口（方法/URL/参数/响应），
       保存到 lookup_api.json —— 之后运行直接复放 API，不再驱动页面
     - 未探测到前，自动分批驱动 UI 查询并捕获响应
  4) 全部数据结构化导出 Excel（openpyxl），嵌套字段以 JSON 字符串保留

用法：
  python sellersprite_deep_dive.py
"""

import asyncio
import json
import os
import re
import sys
import time
from datetime import datetime, timezone, timedelta
from pathlib import Path

from playwright.async_api import async_playwright

BASE = "https://www.sellersprite.com"
HERE = Path(__file__).parent
STATE_FILE = HERE / "storage_state.json"               # 主账号：AI找竞品（可自由更换）
MEMBER_STATE_FILE = HERE / "storage_state_member.json"  # 会员账号：competitor-lookup 专用（固定）
LOOKUP_CFG = HERE / "lookup_api.json"
OUT_DIR = HERE / "deep_results"
HEADLESS = os.environ.get("HEADLESS", "1") != "0"

# ---------- 终端美化 ----------
os.system("")

class C:
    R = "\033[0m"; B = "\033[1m"; D = "\033[2m"
    RED = "\033[31m"; GREEN = "\033[32m"; YEL = "\033[33m"
    BLU = "\033[34m"; MAG = "\033[35m"; CYA = "\033[36m"; ORG = "\033[38;5;208m"

BANNER = f"""{C.ORG}{C.B}
  +--------------------------------------------------+
  |   SellerSprite  Deep-Dive   历史报告 x 竞品详情   |
  +--------------------------------------------------+
{C.R}{C.D}         history -> ASINs -> competitor-lookup -> xlsx{C.R}
"""

def info(m): print(f"{C.BLU}[i]{C.R} {m}")
def ok(m):   print(f"{C.GREEN}[+]{C.R} {m}")
def warn(m): print(f"{C.YEL}[!]{C.R} {m}")
def err(m):  print(f"{C.RED}[x]{C.R} {m}")

# 服务端 createdTime/reportExpiredTime 是 UTC 毫秒时间戳，统一转北京时间显示
CN_TZ = timezone(timedelta(hours=8))
def fmt_cn_time(v):
    """UTC毫秒时间戳 → 北京时间字符串；若已是字符串(如 completedTime)则原样返回。"""
    if isinstance(v, (int, float)) and v > 0:
        return datetime.fromtimestamp(v / 1000, CN_TZ).strftime("%Y-%m-%d %H:%M")
    return str(v) if v else "-"

# ---------- 页面内 fetch 工具（同源携带登录态） ----------
FETCH_GET = "async (u) => { const r = await fetch(u, {headers: {'accept':'application/json'}}); return {s: r.status, j: await r.json().catch(() => null)}; }"
FETCH_POST = """async (x) => {
  const r = await fetch(x.u, {method:'POST', headers:{'content-type':'application/json','accept':'application/json'}, body: x.b});
  return {s: r.status, j: await r.json().catch(() => null)};
}"""

async def api_get(page, path):
    r = await page.evaluate(FETCH_GET, path)
    if r["s"] != 200 or not r["j"]:
        raise RuntimeError(f"GET {path} -> {r['s']}")
    j = r["j"]
    if j.get("code") not in ("OK", 0, None) and not j.get("success", True):
        raise RuntimeError(f"GET {path} -> {j.get('code')} {j.get('message')}")
    return j

async def api_post(page, path, body):
    r = await page.evaluate(FETCH_POST, {"u": path, "b": json.dumps(body)})
    if r["s"] != 200 or not r["j"]:
        raise RuntimeError(f"POST {path} -> {r['s']}")
    return r["j"]

async def ensure_login(pw, state_file=STATE_FILE, label="账号", force=False):
    """force=True 时丢弃旧登录态，弹出浏览器重新登录。"""
    if force and state_file.exists():
        state_file.unlink()
    if state_file.exists():
        ok(f"[{label}] 使用已保存的登录态 {state_file.name}")
        return
    browser = await pw.chromium.launch(headless=False)
    ctx = await browser.new_context()
    pg = await ctx.new_page()
    await pg.goto(BASE)
    input(f"{C.YEL}>>> [{label}] 请在弹出的浏览器中登录后回车…{C.R}")
    await ctx.storage_state(path=str(state_file))
    await browser.close()
    ok(f"[{label}] 登录态已保存到 {state_file.name}")

# ---------- 1) 历史任务 ----------
def _task_list(j):
    d = j.get("data")
    if isinstance(d, list):
        return d
    if isinstance(d, dict):
        # usage-log 的分页包装：data.data 才是任务列表
        inner = d.get("data")
        if isinstance(inner, list):
            return inner
        for k in ("list", "records", "tasks", "rows"):
            if isinstance(d.get(k), list):
                return d[k]
    return []

async def whoami(page, label):
    """调 /v3/api/client 显示当前登录账号，避免搞混双账号。"""
    try:
        j = await api_get(page, "/v3/api/client")
        d = j.get("data") or {}
        lvl = d.get("memberLevel") or "?"
        exp = d.get("expiredDate")
        exp_s = datetime.fromtimestamp(exp / 1000).strftime("%Y-%m-%d") if isinstance(exp, (int, float)) else "-"
        ok(f"[{label}] 当前账号: {C.B}{d.get('nickname') or '-'}{C.R}  "
           f"会员等级={lvl}  到期={exp_s}")
        return d
    except Exception as e:
        warn(f"[{label}] 获取账号信息失败: {e}")
        return {}

async def fetch_history(page, pages=3, page_size=20):
    out, seen = [], set()
    for pn in range(1, pages + 1):
        j = await api_post(page, "/v3/api/ai-analysis/usage-log",
                           {"pageNumber": pn, "pageSize": page_size, "modules": ["CL"]})
        tasks = _task_list(j)
        if not tasks:
            break
        for t in tasks:
            if t.get("taskId") and t["taskId"] not in seen:
                seen.add(t["taskId"])
                out.append(t)
        if len(tasks) < page_size:
            break
    return out

async def dom_history(page):
    """历史页 DOM 兜底：直接解析 history-item 元素（页面按'近7天'等分组，全部组都解析）。"""
    try:
        await page.goto(f"{BASE}/v3/ai-history?historyType=all&backModule=CL",
                        wait_until="domcontentloaded")
        await page.wait_for_timeout(3000)
        htmls = await page.evaluate(
            "() => Array.from(document.querySelectorAll('div.history-item')).map(e => e.outerHTML)")
        out = []
        for h in htmls:
            m = re.search(r"([0-9a-f]{32})", h)               # taskId 是 32 位十六进制
            if not m:
                continue
            tm = re.search(r"分析时间[：:]\s*([0-9\-: ]+)", h)
            asin_m = re.search(r"\(([A-Z0-9]{10})\)", h)
            if "completed" in h:
                st = "COMPLETED"
            elif "分析中" in h or "submitted" in h.lower():
                st = "SUBMITTED"
            else:
                st = "?"
            out.append({"taskId": m.group(1),
                        "createdTime": tm.group(1).strip() if tm else "-",
                        "taskStatus": st,
                        "market": "",
                        "tabTitle": f"ASIN:{asin_m.group(1)}" if asin_m else "-",
                        "_source": "dom"})
        return out
    except Exception as e:
        warn(f"DOM 兜底解析失败（不影响 API 途径）: {e}")
        return []

def pick_task(tasks):
    print(f"\n{C.CYA}历史任务列表{C.R}")
    for i, t in enumerate(tasks, 1):
        st = t.get("taskStatus")
        color = C.GREEN if st == "COMPLETED" else (C.YEL if st == "SUBMITTED" else C.RED)
        m = re.search(r"\(([A-Z0-9]{10})\)", t.get("tabTitle") or "")
        asin = m.group(1) if m else "?"
        print(f"  {C.B}{i:>2}{C.R}. [{color}{st}{C.R}] {fmt_cn_time(t.get('createdTime'))}  "
              f"{t.get('market') or '-'}  {asin}  {C.D}{(t.get('taskId') or '')[:12]}{C.R}")
    while True:
        raw = input(f"{C.CYA}选择任务编号:{C.R} ").strip()
        try:
            idx = int(raw)
            if 1 <= idx <= len(tasks):
                return tasks[idx - 1]
        except ValueError:
            pass
        warn(f"请输入 1-{len(tasks)}")

# ---------- 2) 提取竞品 ASIN ----------
async def extract_asins(page, task_id):
    j = await api_get(page, f"/v3/api/ai-analysis/task/{task_id}")
    d = j.get("data") or {}
    if d.get("taskStatus") != "COMPLETED":
        warn(f"任务状态 {d.get('taskStatus')}，content 可能为空")
    raw = d.get("content") or "[]"
    items = json.loads(raw) if isinstance(raw, str) else raw
    asins, rows = [], []
    for it in items:
        a = (it.get("asin") or "").strip().upper()
        p = (it.get("parent") or "").strip().upper()
        if a:
            asins.append(a)
            rows.append(it)
    # 去重保序
    seen, uniq = set(), []
    for a in asins:
        if a not in seen:
            seen.add(a); uniq.append(a)
    return uniq, rows

# ---------- 3) 竞品查询 ----------
async def probe_lookup(page, asins, chunk):
    """驱动真实页面查询一批，捕获其接口调用；返回 (发现信息, 行数据列表)"""
    page2 = await ctx_new_page(page)
    captured = []

    async def on_response(resp):
        req = resp.request
        try:
            if req.resource_type not in ("xhr", "fetch") and req.method == "GET":
                return
            entry = {"method": req.method, "url": resp.url, "status": resp.status}
            try:
                entry["post_data"] = req.post_data      # 可能为 gzip 二进制，解码失败则置 None
            except Exception:
                entry["post_data"] = None
            try:
                body = await resp.body()
                if len(body) <= 5_000_000:
                    entry["body"] = body.decode("utf-8", "replace")
            except Exception:
                pass
            captured.append(entry)
        except Exception as e:
            warn(f"流量捕获异常（已跳过该响应）: {e}")

    page2.on("response", on_response)
    await page2.goto(f"{BASE}/v3/competitor-lookup", wait_until="domcontentloaded")
    await page2.wait_for_timeout(3000)

    inp = page2.locator("input.el-input__inner[placeholder*='ASIN']").first
    await inp.wait_for(state="visible", timeout=30000)
    await inp.fill(",".join(asins))
    # 按钮定位：页面有多个隐藏的 primary 按钮，必须选可见的那个（优先带"查询"文本）
    btn = page2.locator("button.el-button--primary:visible", has_text="查询")
    if await btn.count() == 0:
        btn = page2.locator("button.el-button--primary:visible", has_text="搜索")
    if await btn.count() == 0:
        btn = page2.locator("button.el-button--primary:visible")
    if await btn.count() == 0:
        raise RuntimeError("页面上没有可见的 primary 查询按钮，请截图发我确认页面结构")
    await btn.first.click()
    info(f"已点击查询按钮（文本={await btn.first.inner_text()!r}），等待响应…")

    rows = []
    for _ in range(60):  # 最多 60s
        await page2.wait_for_timeout(1000)
        done = [c for c in captured
                if "sellersprite.com" in c["url"] and c.get("body", "").strip().startswith("{")]
        if done:
            # 保存本批全部接口响应原始体，供离线分析
            dbg = HERE / "lookup_debug"
            dbg.mkdir(exist_ok=True)
            (dbg / f"probe_{datetime.now():%Y%m%d_%H%M%S}.json").write_text(
                json.dumps({"request_asins": asins,
                            "responses": [{k: v for k, v in c.items() if k != "body"} |
                                          ({} if "body" not in c else
                                           {"body_parsed": json.loads(c["body"])
                                            if c["body"].strip().startswith("{")
                                            else c["body"][:20000]}) for c in done]},
                           ensure_ascii=False), encoding="utf-8")
        if done:
            # 取最后一个含数据的响应
            best = None
            for c in done:
                try:
                    j = json.loads(c["body"])
                except Exception:
                    continue
                if _find_rows(j):
                    best = (c, j)
            if best:
                c, j = best
                rows = _find_rows(j)
                print(f"  [流量] {c['method']} {c['status']} {c['url']}")
                if c.get("post_data"):
                    print(f"         body: {c['post_data'][:300]}")
                break
    await page2.close()

    # 提炼可复放的接口信息
    disc = None
    for c in captured:
        if "sellersprite.com" in c["url"] and "/api/" in c["url"] and c.get("body", "").strip().startswith("{"):
            j = json.loads(c["body"])
            if _find_rows(j):
                u = c["url"].replace(BASE, "")
                disc = {"method": c["method"], "url": u,
                        "post_data": json.loads(c["post_data"]) if c.get("post_data") else None,
                        "discovered_at": datetime.now().isoformat(timespec="seconds")}
                break
    return disc, rows

def _find_rows(j):
    """在响应 JSON 中寻找竞品行列表（list of dict，含 asin 类字段）"""
    def walk(node, depth=0):
        if depth > 6:
            return None
        if isinstance(node, list) and node and isinstance(node[0], dict):
            keys = {k.lower() for k in node[0]}
            if any("asin" in k for k in keys):
                return node
            for x in node:
                r = walk(x, depth + 1)
                if r:
                    return r
        if isinstance(node, dict):
            for v in node.values():
                r = walk(v, depth + 1)
                if r:
                    return r
        return None
    return walk(j)

async def lookup_api(page, cfg, asins, save_raw=True, max_pages=10):
    """
    用已探测的配置直接复放接口；POST 响应带 total/page 分页时自动翻页取全。
    save_raw=True 时把原始响应落盘供分析。max_pages 限制单批最多翻页数。
    """
    fetch_get = ("async (u) => { const r = await fetch(u, "
                 "{headers: {'accept':'application/json'}}); "
                 "return {s: r.status, j: await r.json().catch(() => null)}; }")
    dbg = HERE / "lookup_debug"
    dbg.mkdir(exist_ok=True)
    all_rows, first_j, pages_fetched = [], None, 1
    if cfg["method"] == "GET":
        sep = "&" if "?" in cfg["url"] else "?"
        url = f"{cfg['url']}{sep}asin={','.join(asins)}"
        first_j = (await page.evaluate(fetch_get, url))["j"]
        all_rows = _find_rows(first_j) or []
    else:
        body = json.loads(json.dumps(cfg["post_data"])) if cfg["post_data"] else {}
        _inject_asins(body, asins)
        body["page"] = 1
        j = await api_post(page, cfg["url"], body)
        first_j = j
        d = j.get("data") if isinstance(j, dict) else None
        if not isinstance(d, dict):
            # data 为字符串 = 服务端报错（常见：登陆已失效）
            code = j.get("code") if isinstance(j, dict) else "?"
            msg = (j.get("message") or str(d))[:120] if isinstance(j, dict) else "?"
            if "SESSION" in str(code) or "登陆" in msg or "登录" in msg:
                raise RuntimeError(f"登录态已失效（{msg}）——请在控制台重新登录对应账号后重试")
            raise RuntimeError(f"接口返回异常 {code}: {msg}")
        all_rows = list(d.get("items") or [])
        total = d.get("total") or 0
        size = d.get("size") or body.get("size") or 60
        page_num = 1
        # 翻页：直到取满 total 或达到页数上限（total 缺失时不翻）
        while total and len(all_rows) < min(total, max_pages * size) and page_num < max_pages:
            page_num += 1
            body["page"] = page_num
            j2 = await api_post(page, cfg["url"], body)
            d2 = (j2 or {}).get("data") or {}
            its = d2.get("items") or []
            if not its:
                break
            all_rows.extend(its)
            pages_fetched = page_num
            await asyncio.sleep(1)
    if save_raw and isinstance(first_j, dict):
        (dbg / f"api_{datetime.now():%Y%m%d_%H%M%S}_{asins[0]}.json").write_text(
            json.dumps({"request_asins": asins, "pages_fetched": pages_fetched,
                        "rows_collected": len(all_rows), "response": first_j},
                       ensure_ascii=False), encoding="utf-8")
    return all_rows or (_find_rows(first_j) or [])

def _inject_asins(body, asins):
    def setin(node):
        if isinstance(node, dict):
            for k, v in node.items():
                if isinstance(v, str) and re.fullmatch(r"[A-Z0-9]{10}(,[A-Z0-9]{10})*", v):
                    node[k] = ",".join(asins)
                elif isinstance(v, list) and v and all(isinstance(x, str) and re.fullmatch(r"[A-Z0-9]{10}", x) for x in v):
                    node[k] = asins
                else:
                    setin(v)
        elif isinstance(node, list):
            for x in node:
                setin(x)
    setin(body)

async def ctx_new_page(page):
    return await page.context.new_page()

# ---------- 4) Excel 导出 ----------
def export_excel(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "competitors"
    # 收集全部字段
    keys = []
    for r in rows:
        for k in r:
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFF3E0")
    for r in rows:
        out = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (list, dict)):
                v = json.dumps(v, ensure_ascii=False)
            out.append(v)
        ws.append(out)
    ws.freeze_panes = "A2"
    wb.save(path)
    return len(rows), len(keys)

# ---------- 主流程 ----------
async def main():
    print(BANNER)
    OUT_DIR.mkdir(exist_ok=True)
    async with async_playwright() as pw:
        # ===== 账号分离 =====
        # 主账号（storage_state.json）：AI找竞品/历史任务，可自由更换
        force_main = input(f"{C.YEL}[主账号·AI找竞品] 重新登录（直接回车=是，n 跳过）？{C.R} ").strip().lower() != "n"
        await ensure_login(pw, STATE_FILE, "主账号·AI找竞品", force_main)
        # 会员账号（storage_state_member.json）：competitor-lookup 专用，固定一个
        force_vip = input(f"{C.YEL}[会员账号·竞品查询] 重新登录（回车=沿用已保存，y=重新登录）？{C.R} ").strip().lower() == "y"
        await ensure_login(pw, MEMBER_STATE_FILE, "会员账号·竞品查询", force_vip)

        browser = await pw.chromium.launch(headless=HEADLESS)
        ctx = await browser.new_context(storage_state=str(STATE_FILE))
        ctx.set_default_timeout(30_000)
        page = await ctx.new_page()
        vip_ctx = await browser.new_context(storage_state=str(MEMBER_STATE_FILE))
        vip_ctx.set_default_timeout(30_000)
        vip_page = await vip_ctx.new_page()

        # 主账号：进入同源页面（fetch 需同源 + 登录态）
        await page.goto(f"{BASE}/v3/ai-history?module=CL", wait_until="domcontentloaded")
        await page.wait_for_timeout(3000)
        await whoami(page, "主账号·AI找竞品")
        # 会员账号：进入竞品查询页并验证会员身份
        await vip_page.goto(f"{BASE}/v3/competitor-lookup", wait_until="domcontentloaded")
        await vip_page.wait_for_timeout(3000)
        await whoami(vip_page, "会员账号·竞品查询")
        try:
            j = await api_get(vip_page, "/v3/api/client/pay-status")
            d = j.get("data")
            preview = json.dumps(d, ensure_ascii=False)[:300] if d is not None else json.dumps(j, ensure_ascii=False)[:300]
            info(f"[会员账号] 会员/支付状态: {preview}")
        except Exception as e:
            warn(f"[会员账号] 会员状态获取失败（查询环节可能被拒）: {e}")

        # ---- 1. 历史任务 ----
        info("拉取历史任务列表（API 分页）…")
        tasks = await fetch_history(page, pages=8)
        api_n = len(tasks)
        ok(f"API 返回 {api_n} 个任务")
        # DOM 兜底：历史页元素解析，补齐 API可能漏掉的任务
        dom_tasks = await dom_history(page)
        if dom_tasks:
            known = {t["taskId"] for t in tasks}
            extra = [t for t in dom_tasks if t["taskId"] not in known]
            if extra:
                tasks.extend(extra)
                warn(f"DOM 兜底：页面元素 {len(dom_tasks)} 条 > API {api_n} 条，已补充 {len(extra)} 条")
            else:
                ok(f"DOM 交叉验证：页面元素 {len(dom_tasks)} 条，与 API 一致")
        if not tasks:
            err("未获取到历史任务"); return
        ok(f"合并后共 {len(tasks)} 个任务")
        task = pick_task(tasks)
        tid = task["taskId"]

        # ---- 2. 提取 ASIN ----
        asins, ai_rows = await extract_asins(page, tid)
        ok(f"任务 {tid[:12]} 提取到 {len(asins)} 个竞品 ASIN")
        (OUT_DIR / f"asins_{tid[:12]}.json").write_text(
            json.dumps({"taskId": tid, "asins": asins, "ai_rows": ai_rows},
                       ensure_ascii=False, indent=1), encoding="utf-8")

        # ---- 3. 竞品查询 ----
        cfg = None
        if LOOKUP_CFG.exists():
            cfg = json.loads(LOOKUP_CFG.read_text(encoding="utf-8"))
            ok(f"已加载探测到的查询接口: {cfg['method']} {cfg['url'][:70]}")
        else:
            warn("尚未探测竞品查询接口，将驱动真实页面自动探测（本批会稍慢）")

        chunk = input(f"{C.CYA}每批查询 ASIN 数（回车=20）:{C.R} ").strip()
        chunk = int(chunk) if chunk.isdigit() and int(chunk) > 0 else 20
        if cfg is None and chunk > 20:
            warn(f"首次探测用小批量更稳：已将每批 {chunk} 调整为 20（接口确认后可再用大批量）")
            chunk = 20

        all_rows, meta = [], []
        t0 = time.time()
        for i in range(0, len(asins), chunk):
            batch = asins[i:i + chunk]
            print(f"\n{C.ORG}—— 批次 {i // chunk + 1}/{(len(asins) + chunk - 1) // chunk}: {len(batch)} 个 ASIN ——{C.R}")
            try:
                if cfg:
                    rows = await lookup_api(vip_page, cfg, batch)
                    # 数据行可能不含种子批次信息，补一列
                    for r in rows:
                        r["_batch_asins"] = ",".join(batch)
                else:
                    disc, rows = await probe_lookup(vip_page, batch, chunk)
                    if disc:
                        cfg = disc
                        LOOKUP_CFG.write_text(json.dumps(disc, ensure_ascii=False, indent=1), encoding="utf-8")
                        ok(f"接口已探测并保存到 {LOOKUP_CFG.name}，后续批次将直接复放 API")
                    for r in rows:
                        r["_batch_asins"] = ",".join(batch)
                ok(f"本批返回 {len(rows)} 行")
                all_rows.extend(rows)
            except Exception as e:
                err(f"批次失败: {type(e).__name__}: {e}")
            await asyncio.sleep(2)

        # ---- 4. 导出 ----
        if all_rows:
            fp = OUT_DIR / f"deep_dive_{tid[:12]}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
            n, k = export_excel(all_rows, fp)
            ok(f"导出完成: {fp.name}  ({n} 行 x {k} 列, 用时 {time.time() - t0:.0f}s)")
        else:
            err("未获得任何查询数据——请把终端输出发回来分析接口探测情况")

        await browser.close()

if __name__ == "__main__":
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        print(f"\n{C.YEL}已中断。{C.R}")
