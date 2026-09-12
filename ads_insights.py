# -*- coding: utf-8 -*-
"""
广告洞察模块（ads-insights 关键词历史）
========================================
数据源：/v3/ads-insights 查询 ASIN → 下载的 KeywordHistory-*.xlsx
字段（来自真实样本）：站点/月份/关键词/ABA排名/月搜索量/搜索增长率/月购买量/购买率/
  展示量/点击量/商品数/需供比/点击总占比/转化总占比/货流值/点击前三中位价/
  点击前三评分数/点击前三评分值/PPC价格/所属类目

三层能力：
  1) parse_history_xlsx   解析导出文件 → 关键词×月份 行集
  2) funnel               噪声漏斗：搜索量阈值 + 打分层 + 分层标签
  3) fetch_ads_history    (实验性) 自动化 UI 流：填ASIN→查询→等跳转→触发下载
"""
import asyncio
import glob
import json
import os
import re
from pathlib import Path

import openpyxl

BASE = "https://www.sellersprite.com"
ADS_URL = f"{BASE}/v3/ads-insights"
XDS_STAT = f"{BASE}/v3/api/xds/stat"
HERE = Path(__file__).parent
ADS_DIR = HERE / "adv_reserach"          # 手动/自动下载的 KeywordHistory-*.xlsx 都放这
RES_DIR = HERE / "搭建资源"               # 关键词反查/挖掘等导出样本与资料
ADS_DEBUG = HERE / "ads_debug"

NUM_FIELDS = ("ABA排名", "月搜索量", "搜索增长率", "月购买量", "购买率", "展示量",
              "点击量", "商品数", "需供比", "点击总占比", "转化总占比", "货流值",
              "点击前三评分数", "点击前三评分值")


def _num(v):
    """'2692.0'/'$0.49'/'555701.0' → float；无法解析返回 None"""
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.replace("$", "").replace(",", "").strip()
        try:
            return float(s)
        except ValueError:
            return None
    return None


# ---------- 1) 解析导出文件 ----------

def parse_history_xlsx(fp):
    """解析一个 KeywordHistory-*.xlsx，返回 {market, source, rows:[{keyword, month, 指标...}]}"""
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    allrows = list(ws.iter_rows(values_only=True))
    wb.close()
    hidx = None
    for i, r in enumerate(allrows[:5]):
        if any("关键词" in str(v) for v in r if v):
            hidx = i
            break
    if hidx is None:
        raise ValueError(f"{fp}: 找不到表头行")
    header = [str(h) for h in allrows[hidx]]
    rows = []
    for r in allrows[hidx + 1:]:
        if not any(r):
            continue
        d = dict(zip(header, r))
        kw = str(d.get("关键词") or "").strip()
        if not kw:
            continue
        row = {"keyword": kw, "month": str(d.get("月份") or ""), "market": str(d.get("站点") or "")}
        for f in NUM_FIELDS:
            row[f] = _num(d.get(f))
        row["点击前三中位价"] = d.get("点击前三中位价")
        row["PPC价格"] = d.get("PPC价格")
        row["PPC价格_num"] = _num(d.get("PPC价格"))
        row["所属类目"] = d.get("所属类目")
        rows.append(row)
    # 来源：文件名支持以下形式（三个工具的导出格式统一为"关键词×月份"20字段）
    #   KeywordHistory-<种子词>-<站点>-<日期>.xlsx              （广告洞察）
    #   [ASIN]_KeywordHistory-<种子词>-<站点>-<日期>.xlsx      （广告洞察·手动标注ASIN）
    #   关键词反查-<ASIN>-关键词<词>-<站点>-<日期>.xlsx          （关键词反查：来源=ASIN）
    #   关键词挖掘[_-]<词根>-<站点>-<日期>.xlsx                  （关键词挖掘：来源=词根）
    name = os.path.basename(fp)
    m = re.match(r"(?:([A-Z0-9]{10})[\s_]+-?)?KeywordHistory-(.+)-([A-Z]{2})-\d{8}\.xlsx$", name)
    m_rev = re.match(r"关键词反查-([A-Z0-9]{10})-关键词(.+)-([A-Z]{2})-\d{8}\.xlsx$", name)
    m_miner = re.match(r"关键词挖掘[_-](.+)-([A-Z]{2})-\d{8}\.xlsx$", name)
    source_asin = None
    if m:
        source_asin, source, market = m.group(1), m.group(2), m.group(3)
    elif m_rev:
        source_asin, source, market = m_rev.group(1), f"反查:{m_rev.group(1)}", m_rev.group(3)
    elif m_miner:
        source_asin, source, market = None, m_miner.group(1), m_miner.group(2)
    else:
        source_asin, source, market = None, name, (rows[0]["market"] if rows else "")
    source = source_asin or source   # 有 ASIN 前缀时以竞品 ASIN 作为来源身份
    return {"source": source, "source_asin": source_asin, "market": market,
            "file": name, "rows": rows}


IMPORT_PATTERNS = ("*KeywordHistory-*.xlsx", "关键词反查-*.xlsx", "关键词挖掘*.xlsx")


def import_all(directory=ADS_DIR):
    """把目录下（含搭建资源目录）全部可识别的导出 xlsx 解析合并"""
    out = []
    seen_files = set()
    for d in (directory, RES_DIR):
        for pattern in IMPORT_PATTERNS:
            for fp in sorted(glob.glob(os.path.join(d, pattern))):
                name = os.path.basename(fp)
                if name.startswith("~$") or fp in seen_files:
                    continue
                seen_files.add(fp)
                try:
                    out.append(parse_history_xlsx(fp))
                except Exception as e:
                    out.append({"file": name, "error": str(e), "rows": []})
    return out


# ---------- 2) 噪声漏斗 ----------

def _agg_keyword(month_rows):
    """一个关键词的月份行 → 聚合画像。
    指标取值：按月份倒序逐字段取第一个非空值（快照行最新优先，
    导入文件的官方 ABA/ARA 指标在快照缺该字段时自动补位，不再被覆盖丢失）。"""
    ms = sorted({r["month"] for r in month_rows if r.get("month")}, reverse=True)
    rows = sorted(month_rows, key=lambda r: str(r.get("month") or ""), reverse=True)

    def pick(f):
        for r in rows:
            v = r.get(f)
            if v is not None:
                return v
        return None

    latest = rows[0] if rows else {}
    recent = [r for r in month_rows if r.get("month") in set(ms[:3])]
    return {
        "keyword": pick("keyword"),
        "market": latest.get("market"),
        "latest_month": ms[0] if ms else "-",
        "months_count": len(ms),
        "月搜索量": pick("月搜索量") or 0,
        "搜索增长率": pick("搜索增长率"),
        "近3月搜索量均值": round(sum((r.get("月搜索量") or 0) for r in recent) / max(len(recent), 1)),
        "月购买量": pick("月购买量"),
        "购买率": pick("购买率"),
        "需供比": pick("需供比"),
        "商品数": pick("商品数"),
        "点击量": pick("点击量"),
        "点击总占比": pick("点击总占比"),
        "转化总占比": pick("转化总占比"),
        "ABA排名": pick("ABA排名"),
        "PPC价格": pick("PPC价格"),
        "PPC价格_num": pick("PPC价格_num"),
        "所属类目": pick("所属类目"),
        "trend": {r["month"]: r.get("月搜索量") for r in month_rows if r.get("month")},
    }


def funnel(rows_by_source, min_vol=500, brands=None, enrich_rows=None):
    """
    噪声漏斗。
      rows_by_source: [{source, market, rows:[关键词×月份]}]（多来源=多竞品/多种子）
      brands: 品牌词表（用于识别品牌防御词，如 'anker'）；命中则分层为'品牌词'
      enrich_rows: [关键词×月份行] 指标补位来源（如关键词挖掘官方数据）——
                   逐字段合并进同月行，但**不计入覆盖来源数**（不影响分层/综合分）
    返回 {
      kept: [聚合后的词行 + _sources + _coverage + _score + _tier],
      dropped: n, total_kw: n
    }
    """
    brands = [b.lower() for b in (brands or [])]
    kw_map = {}   # keyword -> {"months": {month: row}, "sources": set}

    def _merge(e, r):
        prev = e["months"].get(r["month"])
        if prev is None:
            e["months"][r["month"]] = dict(r)
        else:
            # 同月多来源（官方快照/抓取估算/趋势行）：逐字段合并，已有非空值优先，
            # 避免"只有搜索量的趋势行"整行覆盖掉含购买率等指标的官方行
            for k, v in r.items():
                if v is not None and prev.get(k) is None:
                    prev[k] = v

    for src in rows_by_source:
        for r in src.get("rows", []):
            kw = r["keyword"]
            e = kw_map.setdefault(kw, {"months": {}, "sources": set(), "market": src.get("market", "")})
            e["sources"].add(src.get("source") or src.get("source_asin") or "?")
            _merge(e, r)
    for r in enrich_rows or []:
        kw = r.get("keyword")
        if not kw:
            continue
        e = kw_map.setdefault(kw, {"months": {}, "sources": set(), "market": ""})
        _merge(e, r)   # 只补指标，不加 sources

    kept, dropped, aggs = 0, 0, []
    import math
    for kw, e in kw_map.items():
        a = _agg_keyword([r for r in e["months"].values()])
        a["keyword"] = kw
        a["market"] = e["market"]
        if (a["月搜索量"] or 0) < min_vol:
            dropped += 1
            continue
        kept += 1
        coverage = len(e["sources"])
        a["_sources"] = sorted(e["sources"])
        a["_coverage"] = coverage
        a["_is_brand"] = any(b in kw.lower() for b in brands)
        # 综合分：对数化搜索量（防品牌巨词支配）× 覆盖热度 × 供给稀缺
        supply = min(a["需供比"] or 1, 10)
        a["_score"] = round(math.log10(1 + (a["月搜索量"] or 0)) * (1 + coverage) * supply, 1)
        aggs.append(a)

    aggs.sort(key=lambda x: (x["_is_brand"], -x["_score"]))   # 品牌词沉底
    for a in aggs:
        if a["_is_brand"]:
            a["_tier"] = "品牌词"
        elif a["_coverage"] >= 3:
            a["_tier"] = "核心词"
        elif a["_coverage"] == 1:
            a["_tier"] = "机会词"
        else:
            a["_tier"] = "共同词"
    return {"kept": aggs, "dropped": dropped, "total_kw": len(kw_map)}


# ---------- 2.5) 广告词价值排序（通用工具，不掺品类经验） ----------
# 判定一律交给 value_top 的价值函数（点击/广告排名/搜索量/bid 四因子）；
# 品牌词标记使用本地库 brand 字段（数据驱动，任何品类通用）。


def _is_brand_word(r):
    """
    品牌词自动检测（数据驱动，无品牌词典、无品类经验）：
    单 token 关键词若在点击Top3商品标题的**前两个token**中以整词出现 ≥2 次 → 品牌词。
    原理：品牌导航词的搜索结果头部必然是品牌方自己的商品，标题以品牌名开头
    （如 UGREEN…/Jackery…/EF ECOFLOW…），而通用词的头部商品标题千差万别。
    """
    kw = str(r.get("keyword") or "").strip().lower()
    if not kw or " " in kw or len(kw) < 3:
        return False
    ara = r.get("araClickTop3") or []
    hits = 0
    for x in ara[:3]:
        t = str(x.get("productTitle") or "").strip().lower()
        if not t:
            continue
        toks = [w.strip(".,-—|") for w in t.split(" ")[:2]]
        if kw in toks:
            hits += 1
    return hits >= 2


def relevance_filter(rows, min_searches=0, brands=None, official=None):
    """
    广告词价值排序与清洗（通用版，价值函数 v2）。
      - 全量行交给 value_top 打价值分并排序（多因子，缺因子按中性值计）；
      - official: {keyword: 官方指标}，参与评分并随结果行输出；
      - min_searches 过滤低量词；
      - 品牌词分层 = 本地库品牌表命中 ∨ 点击Top3标题品牌特征自动检测（数据驱动），
        不剔除任何词。
    返回 (kept, stats)。
    """
    brands = [b.lower() for b in (brands or [])]
    official = official or {}
    rows = [r for r in rows if str(r.get("keyword") or "").strip()]
    kept, stat = [], {"brand": 0, "standard": 0, "low_search": 0}
    scored = value_top(rows, n=len(rows), official=official)
    for score, r in scored:
        kw = str(r.get("keyword") or "").strip()
        searches = r.get("searches") if isinstance(r.get("searches"), (int, float)) else 0
        if searches < min_searches:
            stat["low_search"] += 1
            continue
        is_brand = any(b and b in kw.lower() for b in brands) or _is_brand_word(r)
        o = r.get("_official") or {}
        kept.append({"keyword": kw,
                     "来源ASIN": r.get("_asin") or r.get("asin"),
                     "adId": r.get("adId"),
                     "广告排名": _latest_sp_rank(r),
                     "searches": searches,
                     "clicks": r.get("clicks"),
                     "impressions": r.get("impressions"),
                     "bid": r.get("bid"),
                     "exactPpc": r.get("exactPpc"),
                     "broadPpc": r.get("broadPpc"),
                     "phrasePpc": r.get("phrasePpc"),
                     "cprExact": r.get("cprExact"),
                     "月购买量": o.get("月购买量"),
                     "购买率": o.get("购买率"),
                     "需供比": o.get("需供比"),
                     "价值分": score,
                     "分层": "品牌词" if is_brand else "标准"})
        stat["brand" if is_brand else "standard"] += 1
    kept.sort(key=lambda k: (k["分层"] == "品牌词", -k["价值分"]))   # 品牌词沉底
    return kept, stat


def _latest_sp_rank(r):
    hs = sorted(r.get("spRankHistory") or [], key=lambda h: h.get("date") or "")
    return hs[-1].get("rank") if hs else None


# 价值函数 v2 因子权重（合计=1.0）。全部通用信号、无品类经验；
# 某词缺某因子时按中性值 0.5 计（分母固定为全权重）——不因缺数据被惩罚，
# 也不会像"剩余权重归一"那样让稀疏数据行单因子拉满而霸占榜首；
# 有官方 ARA 数据的词自动获得更完整的证据链，数据越全越能冲高。
VALUE_WEIGHTS = {
    "需求规模": 0.20,      # log(月搜索量)
    "购买体量": 0.20,      # log(官方月购买量) —— 关键词挖掘/导入文件的 ARA 数据
    "购买转化": 0.15,      # 官方购买率（越高的词越接近成交）
    "广告验证·点击": 0.10,  # log(竞品广告点击量) —— 竞品真金白银投出来的
    "广告验证·排名": 0.15,  # 竞品广告排名（越前越好，60名封顶）
    "市场定价": 0.10,      # log(bid/exactPpc) —— 市场对这个词的定价
    "供需空间": 0.10,      # 官方需供比(封顶10) 或 搜索量/商品数
}


def value_top(items, n=10, official=None):
    """
    价值函数 v2：多因子加权打分并取前 n 个。
      official: {keyword(小写): {月购买量/购买率/需供比/...}} 官方指标（关键词挖掘/导入文件）。
      每个因子在批内 min-max 归一（重尾取log，排名取反）；缺因子按中性值 0.5 计（不惩罚、不归一）。
    返回 [(score, item)] 降序。
    """
    import math
    official = official or {}

    def _log(v):
        return math.log10(1 + v) if isinstance(v, (int, float)) and v > 0 else None

    def f_searches(it):
        return _log(it.get("searches"))

    def f_purchases(it):
        return _log((it.get("_official") or {}).get("月购买量"))

    def f_purchase_rate(it):
        v = (it.get("_official") or {}).get("购买率")
        return v if isinstance(v, (int, float)) and v > 0 else None

    def f_clicks(it):
        return _log(it.get("clicks"))

    def f_rank(it):
        r = _latest_sp_rank(it)
        return min(r, 60) if isinstance(r, (int, float)) else None

    def f_bid(it):
        return _log(it.get("bid") or it.get("exactPpc"))

    def f_supply(it):
        o = it.get("_official") or {}
        v = o.get("需供比")
        if not (isinstance(v, (int, float)) and v > 0):
            ap, s = it.get("adProducts"), it.get("searches")
            if isinstance(ap, (int, float)) and ap and isinstance(s, (int, float)) and s:
                v = s / ap
        return min(v, 10) / 10 if isinstance(v, (int, float)) and v > 0 else None

    FACTORS = [
        ("需求规模", f_searches, False),
        ("购买体量", f_purchases, False),
        ("购买转化", f_purchase_rate, False),
        ("广告验证·点击", f_clicks, False),
        ("广告验证·排名", f_rank, True),
        ("市场定价", f_bid, False),
        ("供需空间", f_supply, False),
    ]

    # 预处理：把官方指标挂到行副本上
    items2 = []
    for i, it in enumerate(items):
        o = official.get(str(it.get("keyword") or "").strip().lower())
        items2.append((i, {**it, "_official": o} if o else it))

    # 逐因子批内归一
    norm = {}   # factor -> {i: normalized}
    for name, fn, inv in FACTORS:
        vals = [(i, fn(it)) for i, it in items2]
        vals = [(i, v) for i, v in vals if isinstance(v, (int, float))]
        if not vals:
            continue
        lo = min(v for _, v in vals)
        hi = max(v for _, v in vals)
        for i, v in vals:
            x = 0.5 if hi <= lo else (v - lo) / (hi - lo)
            if inv:
                x = 1 - x
            norm.setdefault(name, {})[i] = x

    scored = []
    for i, it in items2:
        score = 0.0
        for name, _, _ in FACTORS:
            v = norm.get(name, {}).get(i)
            score += VALUE_WEIGHTS[name] * (v if v is not None else 0.5)
        scored.append((round(score, 3), it))
    scored.sort(key=lambda x: -x[0])
    return scored[:n]


def _trends_growth(trends):
    """trends: [{month:'202502', search:n}, …] → (月环比增长率fraction, {'YYYY-MM': search})
    增长率取最近两个有搜索量的月份环比；无足够数据返回 None。"""
    if not isinstance(trends, list):
        return None, {}
    series = []
    for t in trends:
        if not isinstance(t, dict):
            continue
        m = str(t.get("month") or "")
        s = t.get("search") or 0
        if len(m) >= 6 and m[:6].isdigit() and s > 0:
            series.append((f"{m[:4]}-{m[4:6]}", s))
    series.sort()
    growth = None
    if len(series) >= 2 and series[-2][1]:
        growth = round((series[-1][1] - series[-2][1]) / series[-2][1], 4)
    return growth, dict(series)


def full_rows_to_funnel(grouped):
    """
    把 ⚡自动抓取 入库的 ads_kw_full 行（关键词×收割周画像）转换成 funnel 兼容的
    关键词×月份行集，使"②生成广告分析报告"能直接消费自动抓取数据，无需手动放文件。
    快照行映射：month=收割时间戳所在月；searches→月搜索量；clicks→点击量；
      impressions→展示量；bid(缺则exactPpc)→PPC价格；searchesRank→ABA排名；
      adProducts→商品数；searches/adProducts→需供比(估算)；clicks/searches→点击总占比(估算)；
      cvsShareRate→转化总占比；trends 环比→搜索增长率(估算)。
    另按 trends 展开逐月趋势行（搜索量），供"近3月均值"与"搜索量趋势(月)"矩阵使用。
    """
    out = []
    for g in grouped or []:
        rows = []
        for r in g.get("rows", []):
            kw = r.get("keyword")
            if not kw:
                continue
            bid = r.get("bid") or r.get("exactPpc")
            searches = r.get("searches") or 0
            clicks = r.get("clicks") or 0
            ad_prods = r.get("adProducts")
            growth, trend_map = _trends_growth(r.get("trends"))
            rows.append({
                "keyword": kw,
                "month": str(r.get("_taken_at") or "")[:7] or "auto",
                "market": "",
                "月搜索量": searches,
                "点击量": clicks,
                "展示量": r.get("impressions") or 0,
                "PPC价格": bid,
                "PPC价格_num": _num(bid),
                "ABA排名": r.get("searchesRank"),
                "商品数": ad_prods,
                "需供比": round(searches / ad_prods, 1) if searches and ad_prods else None,
                "点击总占比": round(clicks / searches, 4) if searches and clicks else None,
                "转化总占比": r.get("cvsShareRate"),
                "搜索增长率": growth,
            })
            for m, s in trend_map.items():   # 逐月趋势行（只有搜索量）
                rows.append({"keyword": kw, "month": m, "market": "", "月搜索量": s})
        if rows:
            out.append({"source": g.get("source_asin") or "?", "market": "", "rows": rows})
    return out


def parse_kw_miner_items(items):
    """关键词挖掘批量分析接口（POST /v3/api/keyword-miner）的 items → 20字段 关键词×月份行。
    全部为官方值：purchases→月购买量、purchaseRate→购买率、supplyDemandRatio→需供比、
    searchRank→ABA排名、searchesGrowth→搜索增长率（接口为百分数，÷100 归一）。
    同时展开 trends 逐月行（月搜索量），补齐近3月均值与趋势矩阵。"""
    rows = []
    for r in items or []:
        kw = str(r.get("keyword") or "").strip()
        if not kw:
            continue
        month = str(r.get("month") or "").replace(".", "-") or ""   # '2026.08' → '2026-08'
        g = r.get("searchesGrowth")
        clicks, searches = r.get("clicks"), r.get("searches")
        rows.append({
            "keyword": kw, "month": month, "market": "",
            "月搜索量": searches,
            "月购买量": r.get("purchases"),
            "购买率": r.get("purchaseRate"),
            "需供比": r.get("supplyDemandRatio"),
            "商品数": r.get("products"),
            "点击量": clicks,
            "展示量": r.get("impressions"),
            "点击总占比": (round(clicks / searches, 4)
                      if isinstance(clicks, (int, float)) and isinstance(searches, (int, float))
                      and searches else None),   # 官方导出同名列也是该口径
            "转化总占比": r.get("cvsShareRate"),
            "ABA排名": r.get("searchRank"),
            "PPC价格": r.get("bid"),
            "PPC价格_num": _num(r.get("bid")),
            "搜索增长率": (round(g / 100, 4) if isinstance(g, (int, float)) else None),
            "所属类目": "、".join(str(d.get("label") or "")
                           for d in (r.get("departments") or []) if d.get("label")) or None,
        })
        for t in (r.get("trends") or []):
            if not isinstance(t, dict):
                continue
            m = str(t.get("month") or "")
            if len(m) >= 6 and m[:6].isdigit() and (t.get("search") or 0) > 0:
                mm = f"{m[:4]}-{m[4:6]}"
                if mm == month:
                    continue   # 快照行已含该月（且带完整指标），跳过避免同键覆盖
                rows.append({"keyword": kw, "month": mm, "market": "", "月搜索量": t.get("search")})
    return rows


# ---------- 3) 广告组结构 API（xds/stat，已逆向确认） ----------

class SessionExpired(Exception):
    """会员登录态失效（ERR_GLOBAL_SESSION_EXPIRED）"""


def _check_api(j):
    """统一校验：data 变成字符串/报错码 → 抛清晰异常（会话过期单独识别）。"""
    if not isinstance(j, dict):
        raise RuntimeError(f"响应非JSON: {str(j)[:100]}")
    code, msg = j.get("code"), j.get("message") or ""
    data = j.get("data")
    if isinstance(data, str) or (code not in (None, "OK", 0) and not j.get("success", True)):
        text = msg or str(data)[:120]
        if "SESSION" in str(code) or "登陆" in text or "登录" in text:
            raise SessionExpired(f"会员登录态已失效（{text}）——请在控制台重新登录会员账号")
        raise RuntimeError(f"接口错误 {code}: {text}")


async def _api_get(page, url):
    fetch_get = ("async (u) => { const r = await fetch(u, "
                 "{headers: {'accept':'application/json'}}); "
                 "return {s: r.status, j: await r.json().catch(() => null)}; }")
    r = await page.evaluate(fetch_get, url)
    j = (r or {}).get("j")
    _check_api(j)
    return j


async def fetch_ads_groups(page, asin, market_id=1, months=6, interval="week"):
    """
    纯 API 拉取一个 ASIN 的全部广告组（含每周 keywordsCount 时间线）。
    page 需已处于 sellersprite 同源页面（登录态任意）。
    每页原始响应都会落盘 ads_debug/xds_p{n}_<asin>.json 供诊断。
    """
    groups, page_num, meta = [], 1, {}
    dbg = HERE / "ads_debug"
    dbg.mkdir(exist_ok=True)
    while True:
        url = (f"{XDS_STAT}?asin={asin}&market={market_id}&interval={interval}&desc=false"
               f"&pageNum={page_num}&pageSize=50&orderByEffectiveness=false"
               f"&months={months}&combination=false")
        j = await _api_get(page, url)
        try:
            (dbg / f"xds_p{page_num}_{asin}.json").write_text(
                json.dumps({"url": url, "resp": j}, ensure_ascii=False), encoding="utf-8")
        except Exception:
            pass
        d = j.get("data") or {}
        gs = d.get("groups") or []
        groups.extend(gs)
        meta = {"adCount": d.get("adCount"), "campaignCount": d.get("campaignCount"),
                "maxKeywordsCount": d.get("maxKeywordsCount")}
        if not gs:
            break
        if meta["adCount"] and len(groups) >= meta["adCount"]:
            break
        page_num += 1
        await asyncio.sleep(1)
    return {"asin": asin, **meta, "groups": groups}


async def fetch_group_keywords(page, ad_id, begin_date, end_date,
                               market_id=1, interval="week", page_size=50, max_pages=2):
    """某广告组在某周的关键词列表（限 max_pages 页，避免全量翻页触发风控）。"""
    items, page_num = [], 1
    while page_num <= max_pages:
        url = (f"{XDS_STAT}/keyword?adId={ad_id}&interval={interval}"
               f"&beginDate={begin_date}&endDate={end_date}"
               f"&pageSize={page_size}&pageNum={page_num}&keyword="
               f"&orderField=updatedTime&desc=true&market={market_id}")
        j = await _api_get(page, url)
        d = j.get("data") or {}
        its = d.get("items") or []
        items.extend(its)
        total = d.get("total") or 0
        if not its or len(items) >= total:
            break
        page_num += 1
        await asyncio.sleep(1)
    return items


async def fetch_group_history(page, asin, ad_id, market_id=1, interval="week"):
    """某广告组的关键词历史。"""
    url = (f"{XDS_STAT}/history?asin={asin}&market={market_id}"
           f"&interval={interval}&adId={ad_id}")
    j = await _api_get(page, url)
    return j.get("data") or {}


async def ads_full(page, asin, market_id=1, months=6, top_groups=None,
                   weeks_budget=3, log=None):
    """
    一站式：组结构 + 关键词收割（控量版）。
    组按投放热度取前 top_groups；每组按周词数取前 weeks_budget 周，
    每周只取前 2 页（100 词）并用价值函数保留 Top10。
    """
    def _log(m):
        if log:
            log(m)
    r = await fetch_ads_groups(page, asin, market_id, months)
    _log(f"[i] {asin}: 广告组 {r.get('adCount')} / campaign {r.get('campaignCount')} "
         f"/ 关键词总数上限 {r.get('maxKeywordsCount')}")
    groups = r["groups"]
    if top_groups:
        groups = sorted(groups, key=lambda g: -sum((h.get("keywordsCount") or 0)
                                                   for h in (g.get("history") or [])))[:top_groups]
    out = []
    for gi, g in enumerate(groups, 1):
        total = g.get("totalCount") or 0
        hist = sorted(g.get("history") or [], key=lambda h: h.get("beginDate") or 0)
        weeks = [h for h in hist if (h.get("keywordsCount") or 0) > 0]
        if not weeks:
            continue
        weeks = sorted(weeks, key=lambda h: -(h.get("keywordsCount") or 0))[:weeks_budget]
        seen = {}
        for wi, w in enumerate(weeks, 1):
            kws = await fetch_group_keywords(page, g["adId"], w["beginDate"],
                                             w["endDate"], market_id, max_pages=2)
            top = value_top(kws, 10)
            for _, k in top:
                seen.setdefault(str(k.get("id") or k.get("keyword")), k)
            _log(f"    组{gi} 周{wi}/{len(weeks)}（{w.get('keywordsCount')}词/周）"
                 f" 取Top10 累计去重 {len(seen)}/{total}")
            await asyncio.sleep(1)
        out.append({"adId": g.get("adId"), "campaignId": g.get("campaignId"),
                    "组词数": total, "已覆盖词数": len(seen), "keywords": list(seen.values())})
        await asyncio.sleep(1)
    return {"asin": asin, "adCount": r.get("adCount"), "groups_analyzed": len(out),
            "results": out}

def summarize_ads_groups(groups):
    """广告结构画像：每个组的词数/活跃周/连续性/峰值周。"""
    out = []
    for g in groups:
        hist = sorted(g.get("history") or [], key=lambda h: h.get("beginDate") or 0)
        active = [h for h in hist if (h.get("keywordsCount") or 0) > 0]
        peak = max(hist, key=lambda h: h.get("keywordsCount") or 0) if hist else {}
        def _d(ms):
            from datetime import datetime
            return datetime.fromtimestamp(ms / 1000).strftime("%Y-%m-%d") if ms else "-"
        out.append({
            "adId": g.get("adId"), "type": g.get("type"),
            "总词数": g.get("totalCount") or 0,
            "活跃周数": len(active),
            "总周数": len(hist),
            "首投周": _d(active[0].get("beginDate")) if active else "-",
            "末投周": _d(active[-1].get("endDate")) if active else "-",
            "峰值周词数": peak.get("keywordsCount") or 0,
            "峰值周": _d(peak.get("beginDate")),
        })
    out.sort(key=lambda x: -x["总词数"])
    return out


# ---------- 4) UI 自动化（实验性：首跑需观察选择器） ----------

async def fetch_ads_history(page, asin, download_dir=ADS_DIR, timeout_ms=120_000):
    """
    自动化：/v3/ads-insights 填单个 ASIN → 查询 → 等跳转到详情页 → 触发下载。
    返回下载文件路径；失败抛异常并落盘现场。
    选择器集中在此，便于按真实页面调整。
    """
    import asyncio
    download_dir = Path(download_dir)
    download_dir.mkdir(exist_ok=True)
    ADS_DEBUG.mkdir(exist_ok=True)

    await page.goto(ADS_URL, wait_until="domcontentloaded")
    await page.wait_for_timeout(2500)
    inp = page.locator("input.el-input__inner[placeholder*='ASIN']").first
    await inp.wait_for(state="visible", timeout=30000)
    await inp.fill(asin)
    btn = page.locator("button.my-btn").first
    if await btn.count() == 0:
        btn = page.locator("button.el-button--default:visible", has_text="查询")
    await btn.first.click()
    # 等待跳转到详情页 /v3/ads-insights/<id>?q=<ASIN>
    await page.wait_for_url(re.compile(r"/v3/ads-insights/.+\?q="), timeout=timeout_ms)
    await page.wait_for_timeout(4000)

    # 触发下载：悬停 tooltip 链接 → 点击弹出层里的导出项
    tips = page.locator("a.el-tooltip")
    n = await tips.count()
    last_err = None
    for i in range(n):
        try:
            await tips.nth(i).hover()
            await page.wait_for_timeout(600)
            popper = page.locator("div.el-tooltip__popper:visible")
            texts = await popper.all_inner_texts()
            (ADS_DEBUG / f"popper_{asin}_{i}.json").write_text(
                json.dumps({"texts": texts}, ensure_ascii=False), encoding="utf-8")
            target = popper.locator("text=/下载|导出|Export/")
            if await target.count():
                async with page.expect_download(timeout=120_000) as dl:
                    await target.first.click()
                d = await dl.value
                dest = download_dir / d.suggested_filename
                await d.save_as(str(dest))
                return str(dest)
        except Exception as e:
            last_err = e
    raise RuntimeError(f"未能触发下载（尝试了 {n} 个 tooltip）。最后错误: {last_err}；"
                       f"popper 探测结果已存 ads_debug/")
